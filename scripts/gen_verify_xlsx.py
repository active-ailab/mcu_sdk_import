#!/usr/bin/env python3
"""Generate SDK import verification Excel report.

Usage:
  python3 scripts/gen_verify_xlsx.py -o out/Apollo330_AmbiqSuiteC_5.2.0_验证.xlsx \\
    --title "Apollo330更新SDK AmbiqSuiteC_5.2.0验证" \\
    --version AmbiqSuiteC_5.2.0

Fill sheets after generation (or pass --modules-json). Do not invent PASS.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_SIDE = Side(style="thin", color="B4B4B4")
THICK_SIDE = Side(style="medium", color="000000")
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def style_body(ws, start_row: int, end_row: int, cols: int) -> None:
    for r in range(start_row, end_row + 1):
        for c in range(1, cols + 1):
            ws.cell(row=r, column=c).alignment = WRAP


def outline_table(ws, start_row: int, end_row: int, cols: int) -> None:
    """Thin inner grid + thick outer frame around the table range."""
    if end_row < start_row or cols < 1:
        return
    for r in range(start_row, end_row + 1):
        for c in range(1, cols + 1):
            ws.cell(row=r, column=c).border = Border(
                left=THICK_SIDE if c == 1 else THIN_SIDE,
                right=THICK_SIDE if c == cols else THIN_SIDE,
                top=THICK_SIDE if r == start_row else THIN_SIDE,
                bottom=THICK_SIDE if r == end_row else THIN_SIDE,
            )


def set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_info_sheet(wb: Workbook, meta: dict) -> None:
    ws = wb.active
    ws.title = "更新说明"
    ws["A1"] = meta.get("title", "MCU更新SDK验证")
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:B1")

    rows = [
        ("代码提交 / Gerrit topic", meta.get("gerrit", "")),
        ("版本", meta.get("version", "")),
        ("旧版 → 新版", meta.get("version_range", "")),
        ("版本锚点 (VERSION.txt / SHA)", meta.get("version_anchor", "")),
        ("支持芯片", meta.get("chips", "")),
        ("导入日期", meta.get("date", "")),
        ("HM patch 说明", "同名目录覆盖；保留 sdk/ 内 SDK_CHANGE_FOR_HMOS 与 HMI_*，缺失须回补"),
        ("重要说明（升级必读）", meta.get("upgrade_notes", "")),
        ("已知问题", meta.get("known_issues", "")),
        ("分笔提交一览", meta.get("commits", "")),
    ]
    ws["A3"] = "项"
    ws["B3"] = "内容"
    style_header(ws, 3, 2)
    for i, (k, v) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    end_row = 3 + len(rows)
    style_body(ws, 4, end_row, 2)
    outline_table(ws, 3, end_row, 2)
    set_widths(ws, [28, 80])
    ws.row_dimensions[1].height = 24
    for r in range(4, 4 + len(rows)):
        ws.row_dimensions[r].height = 45


def add_module_sheet(wb: Workbook, modules: list[dict]) -> None:
    ws = wb.create_sheet("模块变更验证")
    headers = [
        "模块名称",
        "源码路径",
        "功能描述",
        "SDK修改说明（修改点）",
        "验证记录",
        "验证人员",
        "状态",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))

    # Empty modules → header-only sheet (do not invent placeholder rows).
    for i, m in enumerate(modules, 2):
        ws.cell(row=i, column=1, value=m.get("name", ""))
        ws.cell(row=i, column=2, value=m.get("path", ""))
        ws.cell(row=i, column=3, value=m.get("desc", ""))
        ws.cell(row=i, column=4, value=m.get("changes", ""))
        ws.cell(row=i, column=5, value=m.get("verify", ""))
        ws.cell(row=i, column=6, value=m.get("owner", ""))
        ws.cell(row=i, column=7, value=m.get("status", "未测"))
        ws.row_dimensions[i].height = 60

    end_row = 1 + len(modules) if modules else 1
    if modules:
        style_body(ws, 2, end_row, len(headers))
    outline_table(ws, 1, end_row, len(headers))
    set_widths(ws, [14, 36, 22, 48, 36, 12, 10])


def add_smoke_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("整机冒烟")
    headers = ["测试项", "方法/关注点", "结果", "备注"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))
    items = [
        ("编译通过", "目标产品 defconfig", "未测", ""),
        ("开机启动", "", "未测", ""),
        ("基本显示/触摸", "", "未测", ""),
        ("充电/电池", "", "未测", ""),
        ("休眠唤醒", "", "未测", ""),
    ]
    for i, row in enumerate(items, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=i, column=c, value=v)
    end_row = 1 + len(items)
    style_body(ws, 2, end_row, len(headers))
    outline_table(ws, 1, end_row, len(headers))
    set_widths(ws, [18, 36, 10, 24])


def add_bt_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("蓝牙变更验证")
    headers = ["测试项", "方法/关注点", "结果", "备注"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))
    items = [
        ("广播/连接/断连", "", "未测", "有 BLE 模块变更时填写"),
        ("配对/回连", "", "未测", ""),
        ("数据收发 / 关键 profile", "", "未测", ""),
        ("Radio 固件版本/刷写", "", "未测", ""),
    ]
    for i, row in enumerate(items, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=i, column=c, value=v)
    end_row = 1 + len(items)
    style_body(ws, 2, end_row, len(headers))
    outline_table(ws, 1, end_row, len(headers))
    set_widths(ws, [28, 36, 10, 24])


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("-o", "--output", type=Path, required=True)
    ap.add_argument("--title", default="MCU更新SDK验证")
    ap.add_argument("--version", default="")
    ap.add_argument("--version-range", default="")
    ap.add_argument("--version-anchor", default="")
    ap.add_argument("--chips", default="")
    ap.add_argument("--date", default="")
    ap.add_argument("--gerrit", default="")
    ap.add_argument("--upgrade-notes", default="")
    ap.add_argument("--known-issues", default="")
    ap.add_argument("--commits", default="")
    ap.add_argument(
        "--modules-json",
        type=Path,
        help='JSON list: [{"name","path","desc","changes","verify","owner","status"}]',
    )
    args = ap.parse_args()

    modules: list[dict] = []
    if args.modules_json and args.modules_json.exists():
        modules = json.loads(args.modules_json.read_text(encoding="utf-8"))

    meta = {
        "title": args.title,
        "version": args.version,
        "version_range": args.version_range,
        "version_anchor": args.version_anchor,
        "chips": args.chips,
        "date": args.date,
        "gerrit": args.gerrit,
        "upgrade_notes": args.upgrade_notes,
        "known_issues": args.known_issues,
        "commits": args.commits,
    }

    wb = Workbook()
    add_info_sheet(wb, meta)
    add_module_sheet(wb, modules)
    add_smoke_sheet(wb)
    add_bt_sheet(wb)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"wrote {args.output}")


if __name__ == "__main__":
    main()
